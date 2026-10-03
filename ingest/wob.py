from __future__ import annotations

import datetime as dt
import logging
from collections.abc import Iterator
from pathlib import Path

import httpx
import openpyxl

from app import db
from app.config import DATA_DIR, HTTP_TIMEOUT, USER_AGENT, WOB_HISTORY_URL

log = logging.getLogger(__name__)

SOURCE = "ec_weekly_oil_bulletin"

COUNTRIES = ("DE", "PL", "EU")
PRODUCTS = {"euro95": "e5", "diesel": "diesel"}

SHEETS = (
    ("Prices with taxes", "incl_taxes"),
    ("Prices wo taxes", "excl_taxes"),
)

# Der Cache liegt bewusst neben der Datenbank auf dem PVC. Ein fester Pfad
# unter /app wuerde im Container an readOnlyRootFilesystem scheitern.
CACHE = DATA_DIR / "wob_history.xlsx"


def _product_columns(header: tuple, sheet_basis: str) -> dict[tuple[str, str, str], int]:
    """Map (country, basis, product) -> column index from the header row.

    Die Basis wird bewusst aus dem Blatt abgeleitet, nicht aus dem Spaltennamen:
    "Prices wo taxes" verwendet intern ebenfalls das Wort "tax", sonst
    kollabieren beide Blaetter im Ergebnis und die Nettowerte ueberschreiben
    die Bruttopreise.
    """
    found: dict[tuple[str, str, str], int] = {}
    for idx, cell in enumerate(header):
        if not cell:
            continue
        parts = str(cell).split("_")
        if len(parts) != 5:
            continue
        country, _kind, _basis_word, _tax_word, product = parts
        if country not in COUNTRIES or product not in PRODUCTS:
            continue
        found[(country, sheet_basis, product)] = idx
    return found


def _to_float(raw: object) -> float | None:
    if raw is None or isinstance(raw, str):
        return None
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return None
    if value <= 0:
        return None
    return value


def download(path: Path = CACHE, *, force: bool = False) -> Path:
    if path.exists() and not force:
        log.info("verwende gecachte Datei %s", path)
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    log.info("lade Weekly Oil Bulletin: %s", WOB_HISTORY_URL)
    with httpx.stream(
        "GET", WOB_HISTORY_URL, timeout=HTTP_TIMEOUT, follow_redirects=True,
        headers={"User-Agent": USER_AGENT},
    ) as response:
        response.raise_for_status()
        with path.open("wb") as fh:
            for chunk in response.iter_bytes():
                fh.write(chunk)
    log.info("gespeichert: %s (%d bytes)", path, path.stat().st_size)
    return path


def parse(path: Path = CACHE) -> Iterator[tuple]:
    """Yield (country, fuel, basis, 'weekly', day, value_eur_per_l, unit, currency, sample_size)."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet_name, sheet_basis in SHEETS:
        if sheet_name not in workbook.sheetnames:
            log.warning("Blatt %s fehlt, wird uebersprungen", sheet_name)
            continue
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator)
        columns = _product_columns(header, sheet_basis)
        # Der Bulletin fuehrt die Einheit in Zeile 3 ("1000 l"). Beide Blaetter
        # sind auf 1000Liter umgestellt, wir normieren auf EUR/l.
        if not columns:
            log.warning("keine Produktspalten in %s", sheet_name)
            continue
        for row in iterator:
            raw_day = row[0]
            if not isinstance(raw_day, dt.datetime | dt.date):
                continue
            day = raw_day.date() if isinstance(raw_day, dt.datetime) else raw_day
            for (country, basis_tag, product), col in columns.items():
                value = _to_float(row[col] if col < len(row) else None)
                if value is None:
                    continue
                yield (
                    country,
                    PRODUCTS[product],
                    basis_tag,
                    "weekly",
                    day.isoformat(),
                    round(value / 1000.0, 6),
                    "EUR/l",
                    "EUR",
                    None,
                )


def ingest(*, force_download: bool = False) -> int:
    started = db.utcnow()
    path = download(force=force_download)
    rows = list(parse(path))
    if not rows:
        raise RuntimeError("Weekly Oil Bulletin enthaelt keine verwertbaren Zeilen")
    with db.connect() as conn:
        written = db.upsert_price_points(conn, rows, source=SOURCE, ingested_at=db.utcnow())
        db.record_run(
            conn, SOURCE, started, rows=written, status="ok",
            message=f"parse={len(rows)} aus {path.name}",
        )
    log.info("%s: %d Zeilen geschrieben", SOURCE, written)
    return written
